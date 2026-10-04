"""bb2dash :: ingest/fixtures/extract/make_samples.py

Writes the four synthetic fixtures extract_text.test.mjs reads (brief 100 task 14). Every word is
made up here; nothing comes from course material. Run once, from the repo root, with the locked set:

    uv run --locked --project ingest python ingest/fixtures/extract/make_samples.py

The committed files are the record; a rerun rewrites them with new timestamps inside the zip
containers, which changes their bytes but not their text. Each fixture exercises one rule of
extract_text.py: a blank PDF page and an empty slide or sheet are skipped but keep the numbering of
what follows; docx and pptx tables are joined with " | "; speaker notes get "[notes] ".
"""

from pathlib import Path

import docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent

PDF_PAGES = [
    ["bb2dash extraction fixture", "Page one: synthetic text, no course material."],
    [],
    ["Page three follows a blank page."],
]


def pdf_text_stream(lines):
    """One Helvetica line per entry, top down; an empty list is a blank page."""
    if not lines:
        return b""
    shown = b" T* ".join(b"(" + line.encode("ascii") + b") Tj" for line in lines)
    return b"BT /F1 12 Tf 14 TL 72 720 Td " + shown + b" ET"


def make_pdf(path, pages):
    """A minimal PDF 1.4 by hand: catalog, page tree, one Type1 font, a page and stream per page."""
    body = {}
    next_id = 4
    kids = []
    for lines in pages:
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        stream = pdf_text_stream(lines)
        body[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        body[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id
        )
        kids.append(page_id)
    body[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    body[2] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    body[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    # The binary-marker comment carries a NUL so git sees a binary file and never rewrites its line
    # ends on a Windows checkout (core.autocrlf), which would break the xref byte offsets below.
    out = bytearray(b"%PDF-1.4\n%\x00\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for obj_id in range(1, next_id):
        offsets[obj_id] = len(out)
        out += b"%d 0 obj\n%s\nendobj\n" % (obj_id, body[obj_id])
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % next_id
    for obj_id in range(1, next_id):
        out += b"%010d 00000 n \n" % offsets[obj_id]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (next_id, xref_at)
    path.write_bytes(bytes(out))


def make_docx(path):
    document = docx.Document()
    document.add_heading("bb2dash extraction fixture", level=1)
    document.add_paragraph("A synthetic document, no course material.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph after an empty one.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Week", "Topic"], ["1", "Fixtures"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(path)


def make_pptx(path):
    deck = Presentation()
    title_slide, title_only, blank = deck.slide_layouts[0], deck.slide_layouts[5], deck.slide_layouts[6]

    first = deck.slides.add_slide(title_slide)
    first.shapes.title.text = "bb2dash extraction fixture"
    first.placeholders[1].text = "Synthetic slides, no course material."

    second = deck.slides.add_slide(title_only)
    second.shapes.title.text = "A table slide"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, values in enumerate([["Term", "Meaning"], ["Unit", "One slide"]]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    second.notes_slide.notes_text_frame.text = "Speaker note for slide two."

    deck.slides.add_slide(blank)

    fourth = deck.slides.add_slide(blank)
    fourth.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = (
        "Slide four follows a blank slide."
    )
    deck.save(path)


def make_xlsx(path):
    workbook = openpyxl.Workbook()
    schedule = workbook.active
    schedule.title = "Schedule"
    schedule.append(["Week", "Topic", "Hours"])
    schedule.append([1, "Fixtures", 3.5])
    schedule.append([None, None, None])
    schedule.append([2, "Parity", None])
    schedule.cell(row=3, column=1).value = None  # keep row 3 inside the sheet's range, and empty
    workbook.create_sheet("Empty")
    notes = workbook.create_sheet("Notes")
    notes.append(["Synthetic workbook, no course material."])
    workbook.save(path)


if __name__ == "__main__":
    make_pdf(HERE / "sample.pdf", PDF_PAGES)
    make_docx(HERE / "sample.docx")
    make_pptx(HERE / "sample.pptx")
    make_xlsx(HERE / "sample.xlsx")
    print("wrote sample.pdf, sample.docx, sample.pptx, sample.xlsx")
